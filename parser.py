import docx
import re

def parse_question_bank(file_path):
    """Parses the simple UNIT-based question bank format (question_text, unit, marks)."""
    doc = docx.Document(file_path)
    questions = []
    current_unit = None
    current_marks = None

    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue

        unit_match = re.match(r'UNIT\s*\d+[:\-]?\s*(.*)', text, re.IGNORECASE)
        if unit_match:
            current_unit = text
            continue

        marks_match = re.match(r'^(\d+)\s*Marks?\s*(Questions?)?\s*$', text, re.IGNORECASE)
        if marks_match:
            current_marks = int(marks_match.group(1))
            continue

        question_match = re.match(r'^\d+[\.\)]\s*(.+)', text)
        if question_match:
            questions.append({
                "question": question_match.group(1).strip(),
                "unit": current_unit,
                "marks": current_marks
            })

    return questions


def parse_header(file_path):
    """Extracts header details: code, exam title, subject name, duration, max marks."""
    doc = docx.Document(file_path)
    header = {
        "code": None,
        "exam_title": None,
        "subject_name": None,
        "time_duration": None,
        "max_marks": None
    }
    paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]

    for i, text in enumerate(paragraphs):
        if header["code"] is None:
            m = re.search(r'CODE\s*:?\s*(\S+)', text, re.IGNORECASE)
            if m:
                header["code"] = m.group(1)
                continue

        if header["exam_title"] is None and re.search(r'semester', text, re.IGNORECASE):
            header["exam_title"] = text
            if i + 1 < len(paragraphs):
                header["subject_name"] = paragraphs[i + 1]
            continue

        if header["time_duration"] is None:
            m = re.search(r'Time\s*:\s*([\w\.\s]+?)(?:\t|\s{2,}|Max)', text, re.IGNORECASE)
            if m:
                header["time_duration"] = m.group(1).strip()

        if header["max_marks"] is None:
            m = re.search(r'Max\.?\s*Marks?\s*:?\s*(\d+)', text, re.IGNORECASE)
            if m:
                header["max_marks"] = int(m.group(1))

    return header


def parse_reference_paper(file_path):
    """Parses the real college table-based paper: questions + marks + CO/K codes + OR-alternatives."""
    doc = docx.Document(file_path)
    questions = []
    pending_or = False

    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            row_text = " ".join(cells).strip()

            if not row_text:
                continue

            if row_text.upper() == "OR":
                pending_or = True
                continue

            question_cell = ""
            for c in cells:
                if re.search(r'\(CO\s*\d+,?\s*K\s*\d+\)', c, re.IGNORECASE):
                    question_cell = c
                    break
            if not question_cell:
                question_cell = max(cells, key=len)

            if len(question_cell) < 5:
                continue

            co_k_match = re.search(r'\(CO\s*(\d+),?\s*K\s*(\d+)\)', question_cell, re.IGNORECASE)
            co_number = None
            k_number = None
            if co_k_match:
                co_number = int(co_k_match.group(1))
                k_number = int(co_k_match.group(2))
                question_text = question_cell[:co_k_match.start()].strip()
            else:
                question_text = question_cell.strip()

            marks = None
            for c in cells:
                m = re.match(r'^(\d+)\s*M$', c.strip(), re.IGNORECASE)
                if m:
                    marks = int(m.group(1))
                    break

            if marks is None:
                continue

            k_to_bloom = {
                1: "Remember", 2: "Understand", 3: "Apply",
                4: "Analyze", 5: "Evaluate", 6: "Create"
            }
            bloom_level = k_to_bloom.get(k_number, None)

            questions.append({
                "question": question_text,
                "marks": marks,
                "co_number": co_number,
                "k_level": k_number,
                "bloom_level": bloom_level,
                "is_alternative": pending_or
            })
            pending_or = False

    return questions
def parse_syllabus(file_path):
    doc = docx.Document(file_path)
    units = []
    current_unit = None
    current_topics = ""

    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue

        unit_match = re.match(r'UNIT\s*\d+[:\-]?\s*.*', text, re.IGNORECASE)
        if unit_match:
            if current_unit:
                units.append({"unit_name": current_unit, "topics": current_topics})
            current_unit = text
            current_topics = ""
        elif current_unit:
            current_topics = text

    if current_unit:
        units.append({"unit_name": current_unit, "topics": current_topics})

    return units